import json
import os
import sys
import pytest
from unittest.mock import patch
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

BACKEND_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if BACKEND_DIR not in sys.path:
    sys.path.insert(0, BACKEND_DIR)

from main import app
from db.database import Base, get_db
from db.models import User, Course, Module, Question, AdminLog

SQLALCHEMY_TEST_DATABASE_URL = "sqlite:///:memory:"

engine = create_engine(
    SQLALCHEMY_TEST_DATABASE_URL,
    connect_args={"check_same_thread": False},
    poolclass=StaticPool
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)


def override_get_db():
    db = TestingSessionLocal()
    try:
        yield db
    finally:
        db.close()


@pytest.fixture(autouse=True)
def setup_api_test():
    Base.metadata.create_all(bind=engine)
    app.dependency_overrides[get_db] = override_get_db
    yield
    Base.metadata.drop_all(bind=engine)
    app.dependency_overrides.pop(get_db, None)


@pytest.fixture
def client():
    return TestClient(app)


def test_full_api_flow_and_role_gating(client):
    # 1. Register a student user
    student_signup = client.post("/auth/signup", json={
        "email": "api_student@univ.edu",
        "password": "studentpassword123",
        "role": "student"
    })
    assert student_signup.status_code == 201
    student_token = student_signup.json()["access_token"]
    student_headers = {"Authorization": f"Bearer {student_token}"}

    # 2. Register an admin user
    admin_signup = client.post("/auth/signup", json={
        "email": "api_admin@univ.edu",
        "password": "adminpassword123",
        "role": "admin"
    })
    assert admin_signup.status_code == 201
    admin_token = admin_signup.json()["access_token"]
    admin_headers = {"Authorization": f"Bearer {admin_token}"}

    # 3. ACCEPTANCE CRITERIA: Verify student gets 403 on all /admin/* routes
    res_users = client.get("/admin/users", headers=student_headers)
    assert res_users.status_code == 403

    res_mastery = client.get(f"/admin/users/{student_signup.json()['user']['id']}/mastery", headers=student_headers)
    assert res_mastery.status_code == 403

    res_logs = client.get("/admin/logs", headers=student_headers)
    assert res_logs.status_code == 403

    res_del = client.delete("/admin/courses/fake_course_id", headers=student_headers)
    assert res_del.status_code == 403

    # 4. Student creates a Course
    mock_plan_json = json.dumps({
        "modules": [
            {"title": "Module 1: Intro to Chemistry", "order_index": 1, "topics": ["Atoms", "Molecules"], "prerequisites": []}
        ]
    })
    with patch("agents.planner_agent.generate", return_value=mock_plan_json):
        course_resp = client.post(
            "/courses",
            headers=student_headers,
            json={
                "title": "General Chemistry I",
                "syllabus_raw": "Week 1: Atoms & Molecules\nWeek 2: Periodic Table"
            }
        )
        assert course_resp.status_code == 201
        course_data = course_resp.json()
        course_id = course_data["id"]
        assert course_data["title"] == "General Chemistry I"

    # 5. Student requests next question
    mock_quiz_json = json.dumps({
        "questions": [
            {
                "question_text": "What is the atomic number of Hydrogen?",
                "options_json": ["1", "2", "6", "8"],
                "correct_answer": "1",
                "difficulty": "easy",
                "explanation": "Hydrogen has 1 proton."
            }
        ]
    })
    with patch("agents.examiner_agent.generate", return_value=mock_quiz_json):
        quiz_resp = client.get(f"/courses/{course_id}/next-question", headers=student_headers)
        assert quiz_resp.status_code == 200
        quiz_data = quiz_resp.json()
        assert quiz_data["payload"]["type"] == "question"
        q_item = quiz_data["payload"]["questions"][0]
        q_id = q_item["id"]

    # 6. Student answers question
    mock_eval_json = json.dumps({
        "is_correct": True,
        "feedback": "Correct! Hydrogen has atomic number 1.",
        "key_takeaway": "Atomic number = number of protons."
    })
    with patch("agents.evaluator_agent.generate", return_value=mock_eval_json):
        ans_resp = client.post(
            f"/questions/{q_id}/answer",
            headers=student_headers,
            json={"answer_given": "1", "course_id": course_id}
        )
        assert ans_resp.status_code == 200
        ans_data = ans_resp.json()
        assert ans_data["payload"]["evaluation"]["is_correct"] is True

    # 7. Student uses /courses/{course_id}/explain endpoint
    with patch("api.chat_routes.generate", return_value="Here is a simple explanation of atomic mass..."):
        explain_resp = client.post(
            f"/courses/{course_id}/explain",
            headers=student_headers,
            json={"topic": "Atomic Mass", "student_query": "Why does atomic mass have decimals?"}
        )
        assert explain_resp.status_code == 200
        assert "explanation" in explain_resp.json()

    # 8. Student inspects profile mastery
    profile_resp = client.get("/profile/mastery", headers=student_headers)
    assert profile_resp.status_code == 200
    assert "summary" in profile_resp.json()

    # 9. ACCEPTANCE CRITERIA: Admin actions succeed and produce rows in admin_logs
    db = TestingSessionLocal()
    initial_log_count = db.query(AdminLog).count()

    admin_users_resp = client.get("/admin/users", headers=admin_headers)
    assert admin_users_resp.status_code == 200
    assert len(admin_users_resp.json()) >= 2

    admin_mastery_resp = client.get(f"/admin/users/{student_signup.json()['user']['id']}/mastery", headers=admin_headers)
    assert admin_mastery_resp.status_code == 200

    admin_del_resp = client.delete(f"/admin/courses/{course_id}", headers=admin_headers)
    assert admin_del_resp.status_code == 200

    admin_logs_resp = client.get("/admin/logs", headers=admin_headers)
    assert admin_logs_resp.status_code == 200

    # Verify rows written to admin_logs in database
    final_log_count = db.query(AdminLog).count()
    assert final_log_count >= initial_log_count + 3
    db.close()


def test_student_delete_course_endpoint(client):
    # 1. Register student A and student B
    signup_a = client.post("/auth/signup", json={"email": "student_a@univ.edu", "password": "pass123password", "role": "student"})
    assert signup_a.status_code == 201
    headers_a = {"Authorization": f"Bearer {signup_a.json()['access_token']}"}

    signup_b = client.post("/auth/signup", json={"email": "student_b@univ.edu", "password": "pass123password", "role": "student"})
    assert signup_b.status_code == 201
    headers_b = {"Authorization": f"Bearer {signup_b.json()['access_token']}"}

    # 2. Student A creates a course
    with patch("agents.planner_agent.generate", return_value='{"modules": [{"title": "Mod 1", "order_index": 1, "topics": [], "prerequisites": []}]}'):
        res_course = client.post("/courses", headers=headers_a, json={"title": "Math 101", "syllabus_raw": "Intro to Math"})
        assert res_course.status_code == 201
        course_id = res_course.json()["id"]

    # 3. Student B tries to delete Student A's course -> 403 Forbidden
    res_b_delete = client.delete(f"/courses/{course_id}", headers=headers_b)
    assert res_b_delete.status_code == 403
    assert res_b_delete.json()["detail"] == "Access denied to this course"

    # 4. Student A tries to delete non-existent course -> 404
    res_404 = client.delete("/courses/non-existent-id", headers=headers_a)
    assert res_404.status_code == 404

    # 5. Student A deletes own course -> 200 OK
    res_a_delete = client.delete(f"/courses/{course_id}", headers=headers_a)
    assert res_a_delete.status_code == 200
    assert res_a_delete.json()["course_id"] == course_id

    # 6. Verify course is gone
    get_res = client.get(f"/courses/{course_id}", headers=headers_a)
    assert get_res.status_code == 404

    # Verify courses list is empty for student A
    list_res = client.get("/courses", headers=headers_a)
    assert list_res.status_code == 200
    assert len(list_res.json()) == 0


def test_delete_course_with_full_child_hierarchy(client):
    from db.models import Module, Question, Attempt, MasteryProfile, OrchestratorState, Document, DocumentChunk

    # 1. Register student
    signup = client.post("/auth/signup", json={"email": "hierarchy_student@univ.edu", "password": "pass123password", "role": "student"})
    assert signup.status_code == 201
    user_id = signup.json()["user"]["id"]
    headers = {"Authorization": f"Bearer {signup.json()['access_token']}"}

    # 2. Create Course
    with patch("agents.planner_agent.generate", return_value='{"modules": [{"title": "Bio 1", "order_index": 1, "topics": ["Cells"], "prerequisites": []}]}'):
        res_course = client.post("/courses", headers=headers, json={"title": "Biology 101", "syllabus_raw": "Cell biology"})
        assert res_course.status_code == 201
        course_id = res_course.json()["id"]

    # 3. Populate all child tables directly in DB session to simulate full real-world app state
    db = TestingSessionLocal()
    mod = db.query(Module).filter(Module.course_id == course_id).first()
    assert mod is not None

    q = Question(module_id=mod.id, question_text="What is a cell?", options_json=["A", "B"], correct_answer="A", difficulty="easy")
    db.add(q)
    db.commit()
    db.refresh(q)

    att = Attempt(question_id=q.id, user_id=user_id, answer_given="A", is_correct=True)
    mp = MasteryProfile(module_id=mod.id, user_id=user_id, topic="Cells", score_0to1=0.9, attempts_count=1)
    orch = OrchestratorState(course_id=course_id, user_id=user_id, current_module="Bio 1", status="active")
    doc = Document(course_id=course_id, filename="cells.pdf", storage_url="https://supabase.co/cells.pdf")
    chk = DocumentChunk(course_id=course_id, content="Cells are the basic building blocks", embedding_json=[0.1, 0.2])

    db.add_all([att, mp, orch, doc, chk])
    db.commit()

    # Verify rows exist
    assert db.query(Attempt).filter(Attempt.question_id == q.id).count() == 1
    assert db.query(MasteryProfile).filter(MasteryProfile.module_id == mod.id).count() == 1
    assert db.query(OrchestratorState).filter(OrchestratorState.course_id == course_id).count() == 1
    assert db.query(Document).filter(Document.course_id == course_id).count() == 1
    assert db.query(DocumentChunk).filter(DocumentChunk.course_id == course_id).count() == 1
    db.close()

    # 4. Student deletes course through the API endpoint
    delete_res = client.delete(f"/courses/{course_id}", headers=headers)
    assert delete_res.status_code == 200
    assert delete_res.json()["course_id"] == course_id

    # 5. Verify ALL child records are completely wiped without any foreign key integrity errors
    db2 = TestingSessionLocal()
    assert db2.query(Course).filter(Course.id == course_id).count() == 0
    assert db2.query(Module).filter(Module.course_id == course_id).count() == 0
    assert db2.query(Question).filter(Question.id == q.id).count() == 0
    assert db2.query(Attempt).filter(Attempt.question_id == q.id).count() == 0
    assert db2.query(MasteryProfile).filter(MasteryProfile.module_id == mod.id).count() == 0
    assert db2.query(OrchestratorState).filter(OrchestratorState.course_id == course_id).count() == 0
    assert db2.query(Document).filter(Document.course_id == course_id).count() == 0
    assert db2.query(DocumentChunk).filter(DocumentChunk.course_id == course_id).count() == 0

    # Verify user still exists
    assert db2.query(User).filter(User.id == user_id).count() == 1
    db2.close()


def test_create_course_with_docx_upload_without_manual_text(client):
    import io
    import docx

    signup = client.post("/auth/signup", json={"email": "docx_student@univ.edu", "password": "pass123password", "role": "student"})
    assert signup.status_code == 201
    headers = {"Authorization": f"Bearer {signup.json()['access_token']}"}

    # Generate a real .docx binary in memory
    doc = docx.Document()
    doc.add_paragraph("Unit 1: Quantum Mechanics and Wave Functions")
    doc.add_paragraph("Unit 2: Schrödinger Equation and Potential Wells")
    b = io.BytesIO()
    doc.save(b)
    docx_bytes = b.getvalue()

    mock_plan = '{"modules": [{"title": "Unit 1: Quantum Mechanics", "order_index": 1, "topics": ["Wave Functions"], "prerequisites": []}]}'
    with patch("agents.planner_agent.generate", return_value=mock_plan):
        res = client.post(
            "/courses",
            headers=headers,
            data={"title": "Quantum Physics 301"},
            files={"file": ("quantum_syllabus.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )
        assert res.status_code == 201
        data = res.json()
        assert data["title"] == "Quantum Physics 301"
        assert "Quantum Mechanics" in data["syllabus_raw"]
        assert len(data["modules"]) == 1


def test_create_course_with_image_upload_without_manual_text(client):
    signup = client.post("/auth/signup", json={"email": "img_student@univ.edu", "password": "pass123password", "role": "student"})
    assert signup.status_code == 201
    headers = {"Authorization": f"Bearer {signup.json()['access_token']}"}

    mock_plan = '{"modules": [{"title": "Module 1: Visual Syllabus Overview", "order_index": 1, "topics": ["Diagrams"], "prerequisites": []}]}'
    fake_png = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR"

    with patch("agents.planner_agent.generate", return_value=mock_plan):
        res = client.post(
            "/courses",
            headers=headers,
            data={"title": "Computer Vision 101"},
            files={"file": ("syllabus_chart.png", fake_png, "image/png")}
        )
        assert res.status_code == 201
        data = res.json()
        assert data["title"] == "Computer Vision 101"
        assert len(data["modules"]) == 1


def test_create_course_validation_empty_content(client):
    signup = client.post("/auth/signup", json={"email": "val_student@univ.edu", "password": "pass123password", "role": "student"})
    assert signup.status_code == 201
    headers = {"Authorization": f"Bearer {signup.json()['access_token']}"}

    # Empty title
    res1 = client.post("/courses", headers=headers, json={"title": "", "syllabus_raw": "Some syllabus"})
    assert res1.status_code == 400

    # Empty content and no file
    res2 = client.post("/courses", headers=headers, json={"title": "Valid Title", "syllabus_raw": ""})
    assert res2.status_code == 400



